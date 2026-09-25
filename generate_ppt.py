from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AI-Powered Smart Parking System"
    subtitle.text = "Reducing CO2 Emissions & Traffic Congestion\nUsing IoT, MARL, GNN, & GRU\n\nBTP Project Presentation"

    # Slide 2: Problem Statement
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Problem Statement"
    tf = body_shape.text_frame
    tf.text = "Urban traffic congestion is a growing problem worldwide."
    p = tf.add_paragraph()
    p.text = "Drivers spend an average of 15-20 minutes searching for parking."
    p = tf.add_paragraph()
    p.text = "This leads to increased fuel consumption and heavy CO2 emissions."
    p = tf.add_paragraph()
    p.text = "Existing parking systems lack real-time AI recommendations and environmental tracking."

    # Slide 3: Proposed Solution
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Proposed Solution"
    tf = body_shape.text_frame
    tf.text = "An IoT and AI integrated ecosystem:"
    p = tf.add_paragraph()
    p.text = "Hardware: ESP32 with Ultrasonic sensors to detect slot availability in real-time."
    p = tf.add_paragraph()
    p.text = "Backend AI: Multi-Agent Reinforcement Learning (MARL) for optimal slot allocation."
    p = tf.add_paragraph()
    p.text = "Mobile App: Real-time React Native application with live mapping and CO2 tracking."

    # Slide 4: System Architecture
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "System Architecture"
    tf = body_shape.text_frame
    tf.text = "1. IoT Layer: ESP32 + Ultrasonic Sensors (simulated via Wokwi)."
    p = tf.add_paragraph()
    p.text = "2. Cloud Layer: Firebase Realtime Database."
    p = tf.add_paragraph()
    p.text = "3. AI Backend (Render): Python Flask server runs MARL, GNN, and GRU models continuously."
    p = tf.add_paragraph()
    p.text = "4. User Interface: Expo React Native app with CartoDB interactive map."

    # Slide 5: The AI Core - MARL & GNN
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "AI Core: MARL & GNN"
    tf = body_shape.text_frame
    tf.text = "Multi-Agent Reinforcement Learning (MARL):"
    p = tf.add_paragraph()
    p.text = "Calculates the best slot dynamically based on multiple factors (distance, availability, future predictions)."
    p = tf.add_paragraph()
    p.text = "Graph Neural Networks (GNN):"
    p = tf.add_paragraph()
    p.text = "Models the parking lot as a spatial graph (nodes = slots, edges = roads)."
    p = tf.add_paragraph()
    p.text = "Provides true navigation distances rather than straight-line approximations."

    # Slide 6: The AI Core - GRU
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "AI Core: GRU Predictions"
    tf = body_shape.text_frame
    tf.text = "Gated Recurrent Unit (GRU):"
    p = tf.add_paragraph()
    p.text = "A lightweight Recurrent Neural Network used for time-series prediction."
    p = tf.add_paragraph()
    p.text = "Analyzes historical occupancy data."
    p = tf.add_paragraph()
    p.text = "Predicts if a currently free slot is likely to remain free by the time the driver arrives."
    
    # Slide 7: Environmental Impact
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Environmental Impact Tracking"
    tf = body_shape.text_frame
    tf.text = "Goal: Visibly reduce the carbon footprint of parking."
    p = tf.add_paragraph()
    p.text = "By guiding drivers straight to the optimal slot, the system reduces idle driving time."
    p = tf.add_paragraph()
    p.text = "CO2 Savings = (Time Saved) x (Average Vehicle Emission Rate)."
    p = tf.add_paragraph()
    p.text = "The mobile app features a gamified 'Green Dashboard' showing total CO2 saved in grams."

    # Slide 8: Results & Demo
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Implementation Results"
    tf = body_shape.text_frame
    tf.text = "Latency: Under 3 seconds end-to-end (Hardware -> Cloud -> AI -> App)."
    p = tf.add_paragraph()
    p.text = "Scalability: Firebase and Render setup allows for rapid expansion to hundreds of slots."
    p = tf.add_paragraph()
    p.text = "UI/UX: High-performance React Native Map using free CartoDB tiles."

    # Slide 9: Future Scope
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Future Scope"
    tf = body_shape.text_frame
    tf.text = "Integration with automatic license plate recognition (ALPR)."
    p = tf.add_paragraph()
    p.text = "In-app digital payment gateways (UPI, Credit Cards)."
    p = tf.add_paragraph()
    p.text = "City-wide municipal integration for real-time traffic diversion."
    
    # Slide 10: Conclusion
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Conclusion"
    tf = body_shape.text_frame
    tf.text = "The AI-Powered Smart Parking System successfully bridges IoT hardware with advanced machine learning."
    p = tf.add_paragraph()
    p.text = "It provides a scalable, eco-friendly solution to modern urban traffic congestion."
    p = tf.add_paragraph()
    p.text = ""
    p = tf.add_paragraph()
    p.text = "Thank You!"

    prs.save('Smart_Parking_BTP_Presentation.pptx')
    print("Successfully generated Smart_Parking_BTP_Presentation.pptx")

if __name__ == '__main__':
    create_presentation()
